from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter


PROJECT_ROOT = Path(__file__).resolve().parents[2]

RAW_JOBS_PATH = (
    PROJECT_ROOT
    / "data"
    / "jobs"
    / "raw_jobs.json"
)

EXPORT_PATH = (
    PROJECT_ROOT
    / "data"
    / "exports"
    / "jobpilot_job_tracker.xlsx"
)


def load_jobs() -> list[dict]:
    with open(
        RAW_JOBS_PATH,
        "r",
        encoding="utf-8",
    ) as file:
        payload = json.load(file)

    return payload.get("jobs", [])


def export_jobs() -> None:

    jobs = load_jobs()

    EXPORT_PATH.parent.mkdir(
        parents=True,
        exist_ok=True,
    )

    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Job Opportunities"

    columns = [
        "Raw Job ID",
        "Title",
        "Company",
        "Location",
        "Country",
        "Remote Type",
        "Source",
        "Source Domain",
        "Source URL",
        "Apply URL",
        "Posted Date",
        "Search Task ID",
        "Discovered At",
    ]

    worksheet.append(columns)

    for cell in worksheet[1]:
        cell.font = Font(bold=True)
        cell.fill = PatternFill(
            fill_type="solid",
            fgColor="D9EAF7",
        )

    for job in jobs:

        worksheet.append(
            [
                job.get("raw_job_id"),
                job.get("title"),
                job.get("company"),
                job.get("location"),
                job.get("country"),
                job.get("remote_type"),
                job.get("source"),
                job.get("source_domain"),
                job.get("source_url"),
                job.get("apply_url"),
                job.get("posted_date"),
                job.get("search_task_id"),
                job.get("discovered_at"),
            ]
        )

    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = (
        worksheet.dimensions
    )

    for column_cells in worksheet.columns:

        max_length = 0

        column_letter = get_column_letter(
            column_cells[0].column
        )

        for cell in column_cells:

            value = (
                ""
                if cell.value is None
                else str(cell.value)
            )

            max_length = max(
                max_length,
                len(value),
            )

        worksheet.column_dimensions[
            column_letter
        ].width = min(
            max(max_length + 2, 12),
            50,
        )

    workbook.save(EXPORT_PATH)

    print()
    print("=" * 60)
    print("EXCEL EXPORT COMPLETE")
    print("=" * 60)
    print(
        f"Jobs exported: {len(jobs)}"
    )
    print(
        f"Excel file: {EXPORT_PATH}"
    )


if __name__ == "__main__":
    export_jobs()
