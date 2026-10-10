
from __future__ import annotations

import json
import sys
import tempfile
from pathlib import Path
from typing import Any

import streamlit as st

# ============================================================
# PROJECT PATH
# ============================================================

PROJECT_ROOT = Path(__file__).resolve().parents[1]

if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from app.networking.playwright_setup import ensure_chromium_installed


# ============================================================
# JOB DISCOVERY COMPONENTS
# ============================================================

from app.discovery.search_executor import (
    SearchExecutor,
)
from app.discovery.search_strategy import (
    SearchStrategy,
)
from app.discovery.deduplicator import (
    SourceDeduplicator,
)

from app.discovery.normalizer import (
    normalize_jobs,
    save_jobs,
    OUTPUT_FILE,
)

from app.discovery.job_recency import (
    get_recent_jobs,
)

from app.discovery.deterministic_job_scorer import (
    score_job,
)

from app.storage.job_store import (
    JobStore,
)

from app.storage.database import (
    JobDatabase,
)


# ============================================================
# CONTACT / OUTREACH COMPONENTS
# ============================================================

from app.networking.company_domain_resolver import (
    CompanyDomainResolver,
)

from app.networking.hunter_client import (
    HunterClient,
)

from app.networking.contact_relevance import (
    ContactRelevanceAnalyzer,
)

from app.networking.contact_discovery import (
    ContactDiscovery,
)

from app.storage.contact_store import (
    ContactStore,
)

from app.networking.outreach_generator import (
    OutreachGenerator,
)

from app.networking.outreach_tracker import (
    OutreachTracker,
)

from app.networking.approval_processor import (
    ApprovalProcessor,
)


# ============================================================
# PAGE CONFIG
# ============================================================

st.set_page_config(
    page_title="JobPilot",
    page_icon="🎯",
    layout="wide",
)
# ============================================================
# AUTHENTICATION
# ============================================================

def login_page() -> None:
    st.markdown("""
    <style>
    [data-testid="stSidebar"] {display: none;}
    .stApp {
        background:
            radial-gradient(ellipse at 5% 5%, #102d68 0%, transparent 42%),
            radial-gradient(ellipse at 95% 90%, #0b2451 0%, transparent 38%),
            #050d1d;
        color: #f3f6ff;
    }
    [data-testid="stHeader"] {
        background: transparent;
    }
    .block-container {
        max-width: 1420px;
        padding-top: 3.5rem;
        padding-bottom: 2rem;
    }
    .jp-login-brand {
        font-size: clamp(2.4rem, 4vw, 3.5rem);
        line-height: 1.1;
        font-weight: 800;
        letter-spacing: -1.8px;
        margin-bottom: 0.25rem;
        color: #f8faff;
    }
    .jp-login-brand span {
        color: #3194ff;
    }
    .jp-login-tagline {
        color: #a9b9d6;
        font-size: 1.05rem;
        margin: 0 0 3rem 0.25rem;
    }
    .jp-login-headline {
        font-size: clamp(2.1rem, 3.4vw, 3.15rem);
        line-height: 1.12;
        font-weight: 800;
        letter-spacing: -1.3px;
        margin-bottom: 1.2rem;
    }
    .jp-login-headline span {
        color: #3194ff;
    }
    .jp-login-description {
        color: #b8c6de;
        font-size: 1.08rem;
        line-height: 1.7;
        max-width: 610px;
        margin-bottom: 2rem;
    }
    .jp-feature {
        display: flex;
        align-items: center;
        gap: 1rem;
        margin: 1.15rem 0;
    }
    .jp-feature-icon {
        width: 52px;
        height: 52px;
        flex-shrink: 0;
        display: flex;
        align-items: center;
        justify-content: center;
        border-radius: 14px;
        background: linear-gradient(145deg, #102b54, #142b4c);
        border: 1px solid #28538b;
        color: #55a7ff;
    }
    .jp-feature-icon svg {
        width: 24px;
        height: 24px;
        stroke: currentColor;
        fill: none;
        stroke-width: 1.8;
        stroke-linecap: round;
        stroke-linejoin: round;
    }
    .jp-feature-title {
        color: #f5f7ff;
        font-weight: 700;
        font-size: 1rem;
        margin-bottom: 0.2rem;
    }
    .jp-feature-description {
        color: #9eafca;
        font-size: 0.91rem;
        line-height: 1.45;
    }
    .jp-login-card {
        background: linear-gradient(145deg, #11213b, #0a162a);
        border: 1px solid #233f68;
        border-radius: 22px;
        padding: 2rem 2.1rem 1.5rem;
        box-shadow: 0 24px 70px rgba(0, 0, 0, .25);
    }
    .jp-login-card h2 {
        color: #f5f7ff;
        font-size: 2.1rem;
        font-weight: 800;
        margin: 0 0 .4rem;
        letter-spacing: -.7px;
    }
    .jp-login-card p {
        color: #aabbd5;
        margin-bottom: 1.6rem;
    }
    div[data-testid="stForm"] {
        border: 0;
        padding: 0;
    }
    div[data-testid="stTextInput"] label {
        color: #e4ebfa;
        font-weight: 600;
    }
    div[data-testid="stTextInput"] input {
        background: #15243b;
        color: #f5f7ff;
        border: 1px solid #304767;
        border-radius: 10px;
        min-height: 48px;
    }
    div[data-testid="stTextInput"] input:focus {
        border-color: #3194ff;
        box-shadow: 0 0 0 1px #3194ff;
    }
    div[data-testid="stFormSubmitButton"] button {
        min-height: 50px;
        border-radius: 10px;
        background: linear-gradient(90deg, #2585f5, #3478f6);
        border: 0;
        color: white;
        font-weight: 700;
        margin-top: .7rem;
    }
    div[data-testid="stFormSubmitButton"] button:hover {
        background: #4397ff;
        border: 0;
        color: white;
    }
    .jp-secure-note {
        display: flex;
        align-items: center;
        gap: 12px;
        background: #14253c;
        border: 1px solid #263c5b;
        border-radius: 12px;
        padding: 15px;
        margin-top: 1.5rem;
    }
    .jp-secure-icon {
        font-size: 1.5rem;
    }
    .jp-secure-title {
        color: #58dfaa;
        font-weight: 700;
        margin-bottom: 3px;
    }
    .jp-secure-description {
        color: #9eafca;
        font-size: .87rem;
    }
    .jp-login-footer {
        color: #8295b4;
        font-size: .82rem;
        margin-top: 2rem;
    }
    @media (max-width: 800px) {
        .block-container {
            padding: 1.5rem 1rem;
        }
        .jp-login-tagline {
            margin-bottom: 1.5rem;
        }
        .jp-login-headline {
            font-size: 2rem;
        }
        .jp-login-card {
            padding: 1.5rem 1.2rem;
        }
    }
    </style>
    """, unsafe_allow_html=True)

    left, right = st.columns([1.08, 0.92], gap="large")


    with left:
        st.markdown("""
        <div class="jp-login-brand">
            <span class="jp-compass-mark" aria-label="Career Compass">
                <svg viewBox="0 0 48 48" width="44" height="44"
                     xmlns="http://www.w3.org/2000/svg"
                     style="vertical-align:-5px;margin-right:10px">
                    <defs>
                        <linearGradient id="jpCompassBlue" x1="0" y1="0" x2="1" y2="1">
                            <stop offset="0%" stop-color="#67C3FF"/>
                            <stop offset="100%" stop-color="#2478F5"/>
                        </linearGradient>
                    </defs>
                    <circle cx="24" cy="24" r="21" fill="#102447"
                            stroke="#2E83F7" stroke-width="1.5"/>
                    <circle cx="24" cy="24" r="15" fill="none"
                            stroke="#397DC7" stroke-width="1" opacity=".7"/>
                    <path d="M31.5 13.5L26 26L13.5 34.5L19 22Z"
                          fill="url(#jpCompassBlue)"/>
                    <path d="M31.5 13.5L19 22L26 26Z" fill="#EAF6FF"/>
                    <circle cx="24" cy="24" r="2.3" fill="#FFFFFF"/>
                </svg>
            </span>Job<span>Pilot</span>
        </div>

        <div class="jp-login-tagline">Find. Connect. Outreach. Grow.</div>

        <div class="jp-login-headline">
            Your AI-powered<br>
            <span>Job Search Copilot</span>
        </div>

        <div class="jp-login-description">
            Streamline your job search with intelligent discovery,
            professional contact finding, and personalized outreach
            — all in one place.
        </div>

        <div class="jp-feature">
            <div class="jp-feature-icon">
                <svg viewBox="0 0 24 24" aria-hidden="true">
                    <circle cx="11" cy="11" r="7"></circle>
                    <path d="m16 16 4 4"></path>
                </svg>
            </div>
            <div>
                <div class="jp-feature-title">Discover Opportunities</div>
                <div class="jp-feature-description">
                    Find relevant job openings across global markets.
                </div>
            </div>
        </div>

        <div class="jp-feature">
            <div class="jp-feature-icon">
                <svg viewBox="0 0 24 24" aria-hidden="true">
                    <path d="M16 21v-2a4 4 0 0 0-4-4H8a4 4 0 0 0-4 4v2"></path>
                    <circle cx="10" cy="7" r="4"></circle>
                    <path d="M20 21v-2a4 4 0 0 0-3-3.87"></path>
                    <path d="M16 3.13a4 4 0 0 1 0 7.75"></path>
                </svg>
            </div>
            <div>
                <div class="jp-feature-title">Find Key Contacts</div>
                <div class="jp-feature-description">
                    Identify and connect with relevant decision-makers.
                </div>
            </div>
        </div>

        <div class="jp-feature">
            <div class="jp-feature-icon">
                <svg viewBox="0 0 24 24" aria-hidden="true">
                    <rect x="3" y="5" width="18" height="14" rx="2"></rect>
                    <path d="m3 7 9 6 9-6"></path>
                </svg>
            </div>
            <div>
                <div class="jp-feature-title">Create Personalized Outreach</div>
                <div class="jp-feature-description">
                    Prepare tailored professional outreach drafts.
                </div>
            </div>
        </div>

        <div class="jp-feature">
            <div class="jp-feature-icon">
                <svg viewBox="0 0 24 24" aria-hidden="true">
                    <path d="M3 3v18h18"></path>
                    <path d="M18 17V9"></path>
                    <path d="M13 17V5"></path>
                    <path d="M8 17v-3"></path>
                </svg>
            </div>
            <div>
                <div class="jp-feature-title">Track Your Workflow</div>
                <div class="jp-feature-description">
                    Keep job discovery and outreach organized.
                </div>
            </div>
        </div>

        <div class="jp-login-footer">
            A more focused, organized, and effective job search journey.
        </div>
        """, unsafe_allow_html=True)

    with right:
        st.markdown("""
        <div class="jp-login-card">
            <h2>Welcome Back</h2>
            <p>Sign in to access your JobPilot workspace.</p>
        </div>
        """, unsafe_allow_html=True)

        with st.form("login_form"):
            username = st.text_input(
                "Username",
                placeholder="Enter your username",
            )
            password = st.text_input(
                "Password",
                type="password",
                placeholder="Enter your password",
            )
            submitted = st.form_submit_button(
                "Login  →",
                type="primary",
                use_container_width=True,
            )

            if submitted:
                expected_username = st.secrets["auth"]["username"]
                expected_password = st.secrets["auth"]["password"]

                if (
                    username == expected_username
                    and password == expected_password
                ):
                    st.session_state.authenticated = True
                    st.rerun()
                else:
                    st.error("Invalid username or password.")

        st.markdown("""
        <div class="jp-secure-note">
            <div class="jp-secure-icon">🛡️</div>
            <div>
                <div class="jp-secure-title">Secure Access</div>
                <div class="jp-secure-description">
                    Your workspace is protected by your configured credentials.
                </div>
            </div>
        </div>
        """, unsafe_allow_html=True)


if "authenticated" not in st.session_state:

    st.session_state.authenticated = False


if not st.session_state.authenticated:

    login_page()

    st.stop()
# -----------------------------
# LOGOUT
# -----------------------------

if st.session_state.authenticated:
    if st.sidebar.button(
        "Logout",
        use_container_width=True,
    ):
        st.session_state.authenticated = False
        st.rerun()

# ============================================================
# SESSION STATE
# ============================================================
if "current_cv_name" not in st.session_state:
    st.session_state.current_cv_name = None

if "recent_jobs" not in st.session_state:
    st.session_state.recent_jobs = []

if "search_completed" not in st.session_state:
    st.session_state.search_completed = False

if "contacts" not in st.session_state:
    st.session_state.contacts = {}

if "drafts" not in st.session_state:
    st.session_state.drafts = {}

if "send_results" not in st.session_state:
    st.session_state.send_results = {}

if "search_report" not in st.session_state:
    st.session_state.search_report = {}


# ============================================================
# HELPERS
# ============================================================


def render_profile_value(value, level=0):
    """Render nested CV data as readable Streamlit content."""
    if isinstance(value, dict):
        for key, item in value.items():
            label = str(key).replace("_", " ").replace("-", " ").title()
            if isinstance(item, (dict, list)):
                st.markdown(f"**{label}**")
                render_profile_value(item, level + 1)
            else:
                display = item if item not in (None, "") else "Not available"
                st.markdown(f"- **{label}:** {display}")
    elif isinstance(value, list):
        if not value:
            st.caption("No details available.")
        elif all(not isinstance(item, (dict, list)) for item in value):
            for item in value:
                st.markdown(f"- {item}")
        else:
            for index, item in enumerate(value, start=1):
                if isinstance(item, dict):
                    title = (
                        item.get("name")
                        or item.get("skill")
                        or item.get("title")
                        or f"Item {index}"
                    )
                    with st.container(border=True):
                        st.markdown(f"**{title}**")
                        remaining = {
                            k: v for k, v in item.items()
                            if k not in ("name", "skill", "title")
                        }
                        render_profile_value(remaining, level + 1)
                else:
                    render_profile_value(item, level + 1)
    else:
        st.write(str(value) if value not in (None, "") else "Not available")


def load_candidate_profile() -> dict[str, Any]:
    profile_path = (
        PROJECT_ROOT
        / "data"
        / "candidate"
        / "current_profile.json"
    )

    if not profile_path.exists():
        return {}

    with profile_path.open(
        "r",
        encoding="utf-8",
    ) as file:

        return json.load(file)


def is_ba_title(title: str | None) -> bool:

    value = (
        title or ""
    ).strip().lower()

    signals = (
        "business analyst",
        "senior business analyst",
        "lead business analyst",
        "principal business analyst",
    )

    return any(
        signal in value
        for signal in signals
    )


def get_job_key(job: dict[str, Any]) -> str:

    return (
        f"{job.get('source', '')}"
        f"::"
        f"{job.get('source_job_id', '')}"
    )


def get_contact_name(
    contact: dict[str, Any],
) -> str:

    return (
        contact.get("full_name")
        or contact.get("first_name")
        or "Unknown"
    )


def parse_matched_skills(
    job: dict[str, Any],
) -> list[str]:

    value = job.get(
        "matched_skills"
    )

    if not value:
        return []

    if isinstance(value, list):
        return value

    if isinstance(value, str):

        try:
            parsed = json.loads(value)

            if isinstance(
                parsed,
                list,
            ):
                return parsed

        except json.JSONDecodeError:
            pass

        return [value]

    return []


def format_score(
    score: Any,
) -> str:

    if score is None:
        return "N/A"

    try:
        return str(
            int(float(score))
        )

    except (
        TypeError,
        ValueError,
    ):
        return str(score)


# ============================================================
# JOB SEARCH PIPELINE
# ============================================================


def run_live_job_search() -> dict[str, Any]:

    strategy = SearchStrategy()

    strategy.run()

    executor = SearchExecutor()

    # SearchExecutor.run() returns the list of jobs.
    jobs = executor.run()

    # The execution metrics are persisted separately.
    report_path = (
        PROJECT_ROOT
        / "data"
        / "jobs"
        / "search_execution_report.json"
    )

    if report_path.exists():

        with report_path.open(
            "r",
            encoding="utf-8",
        ) as file:

            execution_report = json.load(file)

    else:

        execution_report = {}

    # --------------------------------------------------------
    # Load raw jobs
    # --------------------------------------------------------

    raw_jobs_path = (
        PROJECT_ROOT
        / "data"
        / "jobs"
        / "raw_jobs.json"
    )

    with raw_jobs_path.open(
        "r",
        encoding="utf-8",
    ) as file:

        raw_payload = json.load(file)

    if isinstance(
        raw_payload,
        dict,
    ):

        raw_jobs = raw_payload.get(
            "jobs",
            [],
        )

    else:

        raw_jobs = raw_payload

    # --------------------------------------------------------
    # Source deduplication
    # --------------------------------------------------------

    deduplicator = SourceDeduplicator()

    unique_jobs, dedup_report = (
        deduplicator.deduplicate(
            raw_jobs
        )
    )

    deduplicated_path = (
        PROJECT_ROOT
        / "data"
        / "jobs"
        / "source_deduplicated_jobs.json"
    )

    with deduplicated_path.open(
        "w",
        encoding="utf-8",
    ) as file:

        json.dump(
            {
                "generated_at": (
                    dedup_report.get(
                        "generated_at"
                    )
                ),
                "total_jobs": len(
                    unique_jobs
                ),
                "jobs": unique_jobs,
            },
            file,
            indent=2,
            default=str,
        )

    # --------------------------------------------------------
    # Normalization
    # --------------------------------------------------------

    normalized_jobs = normalize_jobs(
        unique_jobs
    )

    save_jobs(
        normalized_jobs,
        OUTPUT_FILE,
    )

    # --------------------------------------------------------
    # SQLite
    # --------------------------------------------------------

    job_store = JobStore()

    ingest_result = job_store.ingest()

    # --------------------------------------------------------
    # Recency
    # --------------------------------------------------------

    recent_jobs = get_recent_jobs(
        days=7
    )

    # --------------------------------------------------------
    # Deterministic scoring
    # --------------------------------------------------------

    scored_jobs = []

    for job in recent_jobs:

        result = score_job(
            job
        )

        job_store.update_job_score(
            source=job[
                "source"
            ],
            source_job_id=job[
                "source_job_id"
            ],
            score_result=result,
        )

        enriched_job = dict(
            job
        )

        enriched_job.update(
            result
        )

        scored_jobs.append(
            enriched_job
        )

    job_store.database.connection.commit()

    # --------------------------------------------------------
    # Business Analyst hard inclusion
    # --------------------------------------------------------

    ba_jobs = [
        job
        for job in scored_jobs
        if is_ba_title(job.get("title"))
        and (
            job.get("review_status") or "PENDING"
        ).upper() != "DISMISSED"
    ]

    ba_jobs.sort(
        key=lambda job: (
            job.get(
                "relevance_score"
            )
            or 0
        ),
        reverse=True,
    )

    st.session_state.recent_jobs = (
        ba_jobs
    )

    # --------------------------------------------------------
    # Build UI report
    # --------------------------------------------------------

    tasks_executed = (
        execution_report.get(
            "tasks_selected",
            execution_report.get(
                "tasks_executed",
                0,
            ),
        )
    )

    jobs_discovered = (
        execution_report.get(
            "raw_jobs_found",
            execution_report.get(
                "jobs_discovered",
                len(jobs),
            ),
        )
    )

    st.session_state.search_report = {

        "tasks_executed":
            tasks_executed,

        "jobs_discovered":
            jobs_discovered,

        "new_jobs":
            ingest_result.get(
                "NEW",
                0,
            ),

        "existing_jobs":
            ingest_result.get(
                "EXISTING",
                0,
            ),

        "updated_jobs":
            ingest_result.get(
                "UPDATED",
                0,
            ),

        "recent_jobs":
            len(
                recent_jobs
            ),

        "ba_jobs":
            len(
                ba_jobs
            ),
    }

    st.session_state.search_completed = True

    return (
        st.session_state.search_report
    )

# ============================================================
# CONTACT DISCOVERY
# ============================================================


def discover_contacts(
    job: dict[str, Any],
) -> dict[str, Any]:

    company = (
        job.get("company")
        or ""
    ).strip()

    if not company:

        return {
            "status": "ERROR",
            "message": "Job has no company.",
            "contacts": [],
        }

    # --------------------------------------------------------
    # Resolve company domain
    # --------------------------------------------------------

    resolver = CompanyDomainResolver()

    resolution = resolver.resolve(
        company
    )

    if resolution.status != "VERIFIED":

        resolution = resolver.discover_domain(
            company
        )

    if resolution.status != "VERIFIED":

        return {
            "status": "DOMAIN_UNRESOLVED",
            "message": (
                f"Could not resolve a verified "
                f"domain for {company}."
            ),
            "contacts": [],
        }

    domain = resolution.domain
    # --------------------------------------------------------
    # Hunter + contact relevance
    # --------------------------------------------------------

    database = JobDatabase()

    hunter = HunterClient()

    contact_store = ContactStore(
        database
    )

    relevance = ContactRelevanceAnalyzer()

    discovery = ContactDiscovery(
        hunter_client=hunter,
        contact_store=contact_store,
        relevance_analyzer=relevance,
    )

    matched_skills = parse_matched_skills(
        job
    )

    result = discovery.discover_for_job(
        company=company,
        company_domain=domain,
        canonical_role=(
            job.get(
                "canonical_role"
            )
            or "Business Analyst"
        ),
        matched_skills=matched_skills,
        limit=10,
        max_contacts=5,
        minimum_score=25,
    )

    return result

# ============================================================
# OUTREACH DRAFT
# ============================================================


def generate_draft(
    job: dict[str, Any],
    contact: dict[str, Any],
):

    profile = load_candidate_profile()

    generator = OutreachGenerator(
        candidate_profile=profile
    )

    return generator.generate(
        job=job,
        contact=contact,
    )


# ============================================================
# SEND APPROVED EMAIL
# ============================================================


def approve_and_send(
    run_id: str,
) -> dict[str, Any]:

    tracker = OutreachTracker()

    tracker.update_status(
        run_id,
        "APPROVED",
    )

    processor = ApprovalProcessor(
        tracker=tracker
    )

    return processor.process(
        run_id
    )


# ============================================================
# ============================================================
# JOBPILOT PROFESSIONAL DASHBOARD
# ============================================================
import pandas as pd
from openpyxl import load_workbook

st.markdown("<style>"
":root{--bg:#08111f;--panel:#101d30;--border:#20334e;--muted:#91a4bf}"
".stApp{background:var(--bg);color:#edf4ff}"
"[data-testid='stSidebar']{background:linear-gradient(180deg,#0c1728,#091321);border-right:1px solid var(--border)}"
".block-container{max-width:1800px;padding-top:2.5rem}h1,h2,h3{color:#f4f7ff!important;letter-spacing:-.02em}"
"[data-testid='stMetric']{background:linear-gradient(145deg,#12243a,#0e1b2d);border:1px solid #223752;border-radius:14px;padding:14px 16px;min-height:105px}"
"[data-testid='stMetricLabel']{color:#b8c7dc}[data-testid='stMetricValue']{color:#f7fbff}"
"div[data-testid='stVerticalBlockBorderWrapper']{border-color:#20334e!important;border-radius:14px!important;background:rgba(17,31,50,.72)}"
".stButton>button,.stLinkButton>a{border-radius:9px;font-weight:650;min-height:2.45rem}"
".stButton>button[kind='primary']{background:#2563eb;border:1px solid #4388ff;color:white}"
"[data-testid='stSelectbox']>div>div,[data-testid='stTextInput'] input,[data-testid='stTextArea'] textarea{background:#101d30;border-color:#263b58;border-radius:9px}"
".jp-brand{font-size:1.35rem;font-weight:850;color:#f7fbff}.jp-tagline{font-size:.75rem;color:#8195b1;border-bottom:1px solid #1e2d43;padding-bottom:1rem;margin-bottom:1rem}"
".jp-welcome{font-size:1.8rem;font-weight:800;color:#f4f7ff}.jp-subtitle{color:#9cafc8;margin-bottom:1rem}"
".jp-chip{display:inline-block;background:#1a2c45;border:1px solid #2a405f;color:#bdcde3;border-radius:6px;padding:3px 8px;margin:3px 4px 0 0;font-size:.72rem}"
".jp-score{font-size:1.55rem;font-weight:800;color:#f6f9ff}.jp-progress{height:7px;background:#253750;border-radius:99px;overflow:hidden;margin:6px 0}.jp-progress span{display:block;height:100%;background:linear-gradient(90deg,#22c55e,#4ade80)}hr{border-color:#20334e}"
"</style>",unsafe_allow_html=True)

def _jp_snapshot():
    db=JobDatabase()
    try:
        c=db.connection; jc={r[1] for r in c.execute('PRAGMA table_info(jobs)')}; cc={r[1] for r in c.execute('PRAGMA table_info(contacts)')}
        jobs=c.execute('SELECT COUNT(*) FROM jobs').fetchone()[0] if jc else 0
        pending=c.execute("SELECT COUNT(*) FROM jobs WHERE COALESCE(review_status,'PENDING')='PENDING'").fetchone()[0] if 'review_status' in jc else jobs
        contacts=c.execute('SELECT COUNT(*) FROM contacts').fetchone()[0] if cc else 0
        companies=c.execute("SELECT COUNT(DISTINCT company) FROM contacts WHERE company IS NOT NULL AND TRIM(company)<>''").fetchone()[0] if 'company' in cc else 0
        sources=c.execute("SELECT COALESCE(source,'Unknown'),COUNT(*) FROM jobs GROUP BY source ORDER BY COUNT(*) DESC").fetchall() if 'source' in jc else []
        days=c.execute("SELECT substr(discovered_at,1,10),COUNT(*) FROM jobs WHERE discovered_at IS NOT NULL AND substr(discovered_at,1,10)>=date('now','-13 days') GROUP BY substr(discovered_at,1,10) ORDER BY 1").fetchall() if 'discovered_at' in jc else []
        return {'jobs':jobs,'pending':pending,'contacts':contacts,'companies':companies,'sources':sources,'days':days}
    finally: db.connection.close()

def _jp_outreach():
    path=OutreachTracker().output_path
    if not path.exists(): return []
    try:
        wb=load_workbook(path,read_only=True,data_only=True); rows=list(wb.active.iter_rows(values_only=True)); wb.close()
        if not rows:return []
        heads=[str(v or '') for v in rows[0]]
        return [dict(zip(heads,r)) for r in rows[1:] if any(v is not None for v in r)]
    except Exception:return []

def _jp_counts(rows):
    out={}
    for r in rows:
        s=str(r.get('Status') or 'UNKNOWN').upper();out[s]=out.get(s,0)+1
    return out

def _jp_job_card(job, idx, with_contacts=True):
    key=get_job_key(job); score=job.get('relevance_score')
    try: score=max(0,min(100,int(float(score)))) if score is not None else None
    except (ValueError,TypeError): score=None
    with st.container(border=True):
        a,b,c,d=st.columns([4.3,1.2,1.5,1.25],vertical_alignment='center')
        with a:
            st.markdown('**'+str(job.get('title') or 'Business Analyst')+'**')
            st.write(f"{job.get('company') or 'Company unavailable'} · {job.get('location') or 'Location not specified'}")
            st.caption(f"{job.get('source') or 'Source unavailable'} · {job.get('posted_date') or 'Date unavailable'}")
            skills=parse_matched_skills(job)
            if skills: st.markdown(''.join(f'<span class="jp-chip">{str(x)[:35]}</span>' for x in skills[:5]),unsafe_allow_html=True)
        with b:
            st.caption('MATCH SCORE');st.markdown(f'<div class="jp-score">{score if score is not None else "—"}</div>',unsafe_allow_html=True)
            if score is not None:st.markdown(f'<div class="jp-progress"><span style="width:{score}%"></span></div>',unsafe_allow_html=True)
            st.caption(job.get('relevance_level') or 'Not scored')
        with c:
            st.caption('WORK MODE');st.write(job.get('remote_type') or 'Not specified');st.caption('POSTED');st.write(job.get('posted_date') or 'Date unavailable')
        with d:
            url=job.get('apply_url') or job.get('source_url')
            if url:st.link_button('View Job ↗',url,use_container_width=True)
            if st.button('✕ Dismiss',key=f'jp_dismiss_{key}',use_container_width=True):
                jid=job.get('id')
                if jid is None:st.error('Database ID missing; cannot dismiss this job.')
                else:
                    store=JobStore()
                    try:ok=store.dismiss_job(int(jid))
                    finally:store.database.connection.close()
                    if ok:st.session_state.recent_jobs=[x for x in st.session_state.get('recent_jobs',[]) if get_job_key(x)!=key];st.rerun()
                    else:st.error('Could not dismiss this job.')
        if with_contacts:
            if st.button('Find contacts', key=f'jp_find_{key}'):
                try:
                    ensure_chromium_installed()

                    with st.spinner('Finding relevant contacts...'):
                        st.session_state.contacts[key] = discover_contacts(job)
                except Exception as exc:
                    st.error(f'Contact discovery failed: {exc}')
            result=st.session_state.contacts.get(key)
            if result:
                found_contacts = result.get('contacts', [])
                if not found_contacts:
                    st.warning(
                        result.get('message')
                        or (
                            "No suitable contacts found. "
                            f"Discovery status: {result.get('status', 'UNKNOWN')}"
                        )
                    )
                else:
                    st.success(f"Found {len(found_contacts)} contact(s).")

                for contact in found_contacts:
                    st.markdown(f"**{get_contact_name(contact)}** · {contact.get('position') or 'Position unavailable'} · {contact.get('email') or 'Email unavailable'}")
                    ck=f"{key}::{contact.get('email','')}"
                    if st.button('Generate draft',key=f'jp_draft_{ck}'):
                        try:
                            draft=generate_draft(job,contact); rec=OutreachTracker().record(message=draft,job=job,contact=contact)
                            st.session_state.drafts[ck]={'draft':draft,'run_id':rec['run_id']};st.success(f"Draft created: {rec['run_id']}")
                        except Exception as exc:st.error(f'Draft generation failed: {exc}')
                    dd=st.session_state.drafts.get(ck)
                    if dd:
                        draft=dd['draft'];rid=dd['run_id']
                        with st.expander(f'Draft {rid}',expanded=True):
                            st.text_input('Subject',value=draft.subject,disabled=True,key=f'jp_sub_{rid}')
                            st.text_area('Email body',value=draft.body,disabled=True,key=f'jp_body_{rid}',height=160)
                            if st.button('Approve & Send',type='primary',key=f'jp_send_{rid}'):
                                try:
                                    result=approve_and_send(rid);st.success('Email sent.') if result.get('status')=='SENT' else st.error(str(result))
                                except Exception as exc:st.error(f'Email send failed: {exc}')

with st.sidebar:
    st.markdown('<div class="jp-brand">🎯 JobPilot</div>',unsafe_allow_html=True)
    st.markdown('<div class="jp-tagline">Find. Connect. Outreach. Grow.</div>',unsafe_allow_html=True)
    page=st.radio('NAVIGATION',['Dashboard','Job Discovery','Contacts','Outreach','Follow-ups','Candidate Profile'],label_visibility='collapsed',key='jp_navigation')
    st.divider();cv=PROJECT_ROOT/'data'/'candidate'/'cv'/'current_cv.pdf';st.markdown('**Current CV**');st.caption(cv.name if cv.exists() else 'No current CV file found')

snap=_jp_snapshot(); outreach=_jp_outreach(); statuses=_jp_counts(outreach)
if page=='Dashboard':
    st.markdown('<div class="jp-welcome">Welcome back, Prashil! 👋</div>',unsafe_allow_html=True)
    st.markdown('<div class="jp-subtitle">Track opportunities, connect with the right people, and grow your network.</div>',unsafe_allow_html=True)
    metrics=[('Jobs Discovered',snap['jobs'],'Persisted in SQLite'),('Pending Review',snap['pending'],'Not dismissed'),('Contacts Found',snap['contacts'],f"Across {snap['companies']} companies"),('Outreach Drafts',statuses.get('DRAFT',0),'Awaiting approval'),('Emails Sent',statuses.get('SENT',0),'Confirmed sent status'),('Follow-ups Due','—','Follow-up dates not tracked yet')]
    for col,(label,value,note) in zip(st.columns(6),metrics):
        with col:st.metric(label,value);st.caption(note)
    x,y,z=st.columns([1.5,1.15,1.05])
    with x:
        with st.container(border=True):
            st.markdown('**Job Discovery Trend**');st.caption('Jobs recorded by discovery date · last 14 days')
            if snap['days']:st.line_chart(pd.DataFrame(snap['days'],columns=['Date','Jobs']).set_index('Date'),height=220)
            else:st.info('No dated discovery history available.')
    with y:
        with st.container(border=True):
            st.markdown('**Outreach Status**');st.caption('Records in the outreach tracker');data={a:b for a,b in statuses.items() if b}
            if data:st.bar_chart(pd.DataFrame({'Records':data}),height=220)
            else:st.info('No outreach records found.')
    with z:
        with st.container(border=True):
            st.markdown('**Top Job Sources**');st.caption('Stored opportunities by source')
            if snap['sources']:st.bar_chart(pd.DataFrame(snap['sources'],columns=['Source','Jobs']).set_index('Source'),height=220)
            else:st.info('No source data available.')
    with st.container(border=True):
        l,r=st.columns([4,1.5],vertical_alignment='center')
        with l:st.markdown('## Live Job Discovery');st.caption('Search live sources, update SQLite, and score opportunities against your profile.')
        with r:
            if st.button('🔍 Run Live Job Search',type='primary',use_container_width=True,key='jp_dash_search'):
                try:
                    with st.spinner('Running live job discovery...'):st.session_state.search_report=run_live_job_search()
                    st.rerun()
                except Exception as exc:st.error(f'Job search failed: {exc}')
        current=st.session_state.get('recent_jobs',[])
        if current:
            st.caption(f'{len(current)} eligible Business Analyst jobs in current search results')
            for i,job in enumerate(current[:3]):_jp_job_card(job,i,False)
            if st.button('Open all job results →'):st.session_state.jp_navigation='Job Discovery';st.rerun()
        else:st.info('Run a live search to populate this area with current opportunities.')
elif page=='Job Discovery':
    st.markdown('<div class="jp-welcome">Live Job Discovery</div>',unsafe_allow_html=True);st.caption('Review opportunities, inspect match scores, and find relevant contacts.')
    a,b=st.columns([4,1.4])
    with a:q=st.text_input('Search by job title, company or keywords',placeholder='Business Analyst, transformation, insurance')
    with b:
        if st.button('Run Live Job Search',type='primary',use_container_width=True,key='jp_jobs_search'):
            try:
                with st.spinner('Searching live sources...'):run_live_job_search()
                st.rerun()
            except Exception as exc:st.error(f'Job search failed: {exc}')
    jobs=st.session_state.get('recent_jobs',[])
    if q:
        term=q.lower();jobs=[j for j in jobs if term in str(j.get('title','')).lower() or term in str(j.get('company','')).lower() or term in str(j.get('description','')).lower()]
    st.caption(f'Showing up to 10 of {len(jobs)} eligible jobs')
    if not jobs:st.info('Run Live Job Search to load current results.')
    for i,job in enumerate(jobs[:10]):_jp_job_card(job,i,True)
elif page=='Contacts':
    st.markdown('<div class="jp-welcome">Contacts & Networking</div>',unsafe_allow_html=True);st.caption('Persisted contacts. Contact-to-job relationships are not stored separately yet.')
    db=JobDatabase()
    try:contacts=[dict(r) for r in db.connection.execute('SELECT * FROM contacts ORDER BY updated_at DESC').fetchall()]
    except Exception as exc:contacts=[];st.error(f'Could not load contacts: {exc}')
    finally:db.connection.close()
    if contacts:
        df=pd.DataFrame(contacts);q=st.text_input('Filter contacts',placeholder='Name, company, role or email')
        if q:
            mask=df.astype(str).apply(lambda col:col.str.contains(q,case=False,na=False)).any(axis=1);df=df[mask]
        cols=[c for c in ['first_name','last_name','position','company','email','email_verification_status','email_confidence','linkedin_url'] if c in df.columns];st.dataframe(df[cols],use_container_width=True,hide_index=True)
    else:st.info('No saved contacts yet. Discover contacts from a job card.')
elif page=='Outreach':
    st.markdown('<div class="jp-welcome">Outreach Centre</div>',unsafe_allow_html=True);st.caption('Drafts and send outcomes. Emails require explicit approval.')
    for col,s in zip(st.columns(5),['DRAFT','APPROVED','SENT','SEND_FAILED','SENDING']):
        with col:st.metric(s.replace('_',' ').title(),statuses.get(s,0))
    if outreach:
        df=pd.DataFrame(outreach);cols=[c for c in ['Run ID','Company','Role','Contact Name','Contact Email','Subject','Status','Created At','Sent At','Error Message'] if c in df.columns];st.dataframe(df[cols],use_container_width=True,hide_index=True)
        ids=[str(v) for v in df['Run ID'].dropna()] if 'Run ID' in df else [];selected=st.selectbox('Select an outreach record',['']+ids)
        if selected:
            rec=next((r for r in outreach if str(r.get('Run ID'))==selected),{})
            with st.expander(f'Details — {selected}',expanded=True):
                st.write(f"**{rec.get('Company','')}** · {rec.get('Role','')}");st.write(f"**To:** {rec.get('Contact Name','')} · {rec.get('Contact Email','')}");st.write(f"**Subject:** {rec.get('Subject','')}")
                st.text_area('Email body',value=str(rec.get('Email Body') or ''),disabled=True,key=f'jp_detail_{selected}',height=180)
                if str(rec.get('Status','')).upper() in ('DRAFT','SEND_FAILED') and st.button('Approve & Send selected email',type='primary',key=f'jp_approve_{selected}'):
                    try:
                        result=approve_and_send(selected);st.success('Email sent.') if result.get('status')=='SENT' else st.error(str(result));st.rerun()
                    except Exception as exc:st.error(f'Could not send: {exc}')
    else:st.info('No outreach records found. Generate a draft from a job contact panel.')
elif page=='Follow-ups':
    st.markdown('<div class="jp-welcome">Follow-ups</div>',unsafe_allow_html=True);st.info('The current tracker does not store follow-up dates or replies, so JobPilot cannot calculate a real due count yet.')
elif page=='Candidate Profile':
    st.markdown('<div class="jp-welcome">Candidate Profile</div>',unsafe_allow_html=True);st.caption('Upload and analyze the CV used for job matching and outreach.')
    uploaded=st.file_uploader('Upload CV',type=['pdf','docx'],key='jp_cv_upload')
    if uploaded is not None and st.button('Analyze CV',type='primary',key='jp_analyze_cv'):
        temp=None
        try:
            with tempfile.NamedTemporaryFile(delete=False,suffix=Path(uploaded.name).suffix.lower()) as f:f.write(uploaded.getbuffer());temp=Path(f.name)
            with st.spinner('Analyzing CV...'):profile=CVIngestionService().ingest(temp)
            st.success('CV analyzed successfully.')
            st.session_state['jp_cv_analysis_message'] = 'CV analyzed successfully.'
        except Exception as exc:st.error(f'CV analysis failed: {exc}')
        finally:
            if temp is not None and temp.exists():temp.unlink()
    profile=load_candidate_profile()
    if profile:
        st.markdown('### Professional Summary')
        summary = profile.get('professional_summary', '')
        if summary:
            st.write(summary)
        else:
            st.caption('No professional summary was extracted.')

        st.markdown('### Core Skills')
        capabilities = profile.get('capabilities', {})
        skills = capabilities.get('skills', []) if isinstance(capabilities, dict) else []
        if skills:
            st.markdown(' '.join(f'`{skill}`' for skill in skills))
        else:
            st.caption('No skills were extracted.')

        st.markdown('### Profile Details')
        details = {
            key.replace('_', ' ').title(): value
            for key, value in profile.items()
            if key not in ('professional_summary', 'capabilities', 'source_cv')
        }
        if profile.get('source_cv'):
            st.caption(f"Source CV: {profile['source_cv']}")

        if details:
            for label, value in details.items():
                st.markdown(f"### {label}")
                render_profile_value(value)

        other_capabilities = {
            key: value for key, value in capabilities.items()
            if key != 'skills'
        } if isinstance(capabilities, dict) else {}
        if other_capabilities:
            st.markdown('### Additional Capabilities')
            for label, value in other_capabilities.items():
                st.markdown(f"#### {label.replace('_', ' ').title()}")
                render_profile_value(value)
    else:
        st.info('No extracted candidate profile found.')

