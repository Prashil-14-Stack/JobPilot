from __future__ import annotations

import time

from openpyxl import load_workbook

from app.networking.approval_processor import process_approval
from app.networking.outreach_tracker import OutreachTracker


class ApprovalWatcher:
    """
    Watches the outreach tracker for processable records.

    DRAFT
        ↓
    APPROVED
        ↓
    ApprovalWatcher detects it
        ↓
    ApprovalProcessor sends it
        ↓
    SENT

    If sending fails:

    SEND_FAILED
        ↓
    ApprovalWatcher detects it
        ↓
    ApprovalProcessor retries it
        ↓
    SENT
    """

    POLL_INTERVAL_SECONDS = 30

    def __init__(
        self,
        tracker: OutreachTracker | None = None,
    ) -> None:

        self.tracker = (
            tracker
            or OutreachTracker()
        )

    def find_processable_records(
        self,
    ) -> list[str]:

        if not self.tracker.output_path.exists():
            return []

        workbook = load_workbook(
            self.tracker.output_path,
            read_only=True,
        )

        worksheet = workbook.active

        headers = [
            cell.value
            for cell in worksheet[1]
        ]

        try:
            run_id_column = headers.index(
                "Run ID"
            )

            status_column = headers.index(
                "Status"
            )

        except ValueError:
            workbook.close()
            return []

        processable = []

        for row in worksheet.iter_rows(
            min_row=2,
            values_only=True,
        ):

            if len(row) <= max(
                run_id_column,
                status_column,
            ):
                continue

            run_id = row[
                run_id_column
            ]

            status = row[
                status_column
            ]

            if run_id and str(status).upper() in (
                "APPROVED",
                "SEND_FAILED",
            ):
                processable.append(
                    str(run_id)
                )

        workbook.close()

        return processable

    def process_once(self) -> list[dict]:

        processable_records = (
            self.find_processable_records()
        )

        results = []

        for run_id in processable_records:

            try:

                result = process_approval(
                    run_id
                )

                results.append(
                    result
                )

            except Exception as exc:

                results.append(
                    {
                        "run_id": run_id,
                        "status": "ERROR",
                        "action": "FAILED",
                        "error": str(exc),
                    }
                )

        return results

    def run(self) -> None:

        print(
            "JobPilot approval watcher started."
        )

        print(
            f"Polling every "
            f"{self.POLL_INTERVAL_SECONDS} seconds."
        )

        print(
            "Press Ctrl+C to stop."
        )

        try:

            while True:

                results = (
                    self.process_once()
                )

                for result in results:

                    print(
                        result
                    )

                time.sleep(
                    self.POLL_INTERVAL_SECONDS
                )

        except KeyboardInterrupt:

            print(
                "\nJobPilot approval watcher stopped."
            )


if __name__ == "__main__":

    watcher = ApprovalWatcher()

    watcher.run()