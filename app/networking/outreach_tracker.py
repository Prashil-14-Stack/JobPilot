from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook


class OutreachTracker:
    """
    Maintains the persistent JobPilot outreach Excel tracker.

    SQLite remains the source of truth.
    Excel is the human-readable outreach history/export.
    """

    OUTPUT_PATH = Path(
        "data/exports/jobpilot_outreach_tracker.xlsx"
    )

    HEADERS = [
        "Run ID",
        "Job ID",
        "Company",
        "Role",
        "Location",
        "Job Source",
        "Contact Name",
        "Contact Email",
        "Contact Position",
        "Contact Function",
        "Subject",
        "Email Body",
        "CV Path",
        "Status",
        "Created At",
        "Sent At",
        "Gmail Message ID",
        "Error Message",
        "Retry Count",
    ]

    def __init__(
        self,
        output_path: Path | None = None,
    ) -> None:

        self.output_path = (
            output_path
            or self.OUTPUT_PATH
        )
    def _ensure_columns(
        self,
        worksheet,
    ) -> None:

        existing_headers = [
            cell.value
            for cell in worksheet[1]
        ]

        for header in self.HEADERS:

            if header not in existing_headers:

                worksheet.cell(
                    row=1,
                    column=worksheet.max_column + 1,
                ).value = header

                existing_headers.append(
                    header
                )
    def ensure_schema(self) -> None:

        if not self.output_path.exists():
            return

        workbook = load_workbook(
            self.output_path
        )

        worksheet = workbook.active

        self._ensure_columns(
            worksheet
        )

        workbook.save(
            self.output_path
        )

    def mark_sending(
        self,
        run_id: str,
    ) -> dict[str, Any]:

        if not self.output_path.exists():
            raise FileNotFoundError(
                f"Outreach tracker not found: "
                f"{self.output_path}"
            )

        workbook = load_workbook(
            self.output_path
        )

        worksheet = workbook.active

        headers = [
            cell.value
            for cell in worksheet[1]
        ]

        run_id_column = (
            headers.index("Run ID") + 1
        )

        status_column = (
            headers.index("Status") + 1
        )

        for row_number in range(
            2,
            worksheet.max_row + 1,
        ):

            current_run_id = worksheet.cell(
                row=row_number,
                column=run_id_column,
            ).value

            if current_run_id != run_id:
                continue

            current_status = (
                worksheet.cell(
                    row=row_number,
                    column=status_column,
                ).value
                or ""
            ).upper()

            if current_status != "APPROVED":
                raise PermissionError(
                    f"Outreach {run_id} must be "
                    f"APPROVED before sending. "
                    f"Current status: {current_status}"
                )

            worksheet.cell(
                row=row_number,
                column=status_column,
            ).value = "SENDING"

            workbook.save(
                self.output_path
            )

            return {
                "run_id": run_id,
                "previous_status": current_status,
                "status": "SENDING",
            }

        raise ValueError(
            f"Outreach record not found: {run_id}"
        )

    def mark_sent(
        self,
        run_id: str,
        gmail_message_id: str,
    ) -> dict[str, Any]:

        if not self.output_path.exists():
            raise FileNotFoundError(
                f"Outreach tracker not found: "
                f"{self.output_path}"
            )

        workbook = load_workbook(
            self.output_path
        )

        worksheet = workbook.active

        headers = [
            cell.value
            for cell in worksheet[1]
        ]

        run_id_column = (
            headers.index("Run ID") + 1
        )

        status_column = (
            headers.index("Status") + 1
        )

        sent_at_column = (
            headers.index("Sent At") + 1
        )

        message_id_column = (
            headers.index("Gmail Message ID") + 1
        )

        error_column = (
            headers.index("Error Message") + 1
        )
        retry_count_column = (
            headers.index("Retry Count") + 1
        )
        for row_number in range(
            2,
            worksheet.max_row + 1,
        ):

            current_run_id = worksheet.cell(
                row=row_number,
                column=run_id_column,
            ).value

            if current_run_id != run_id:
                continue

            current_status = (
                worksheet.cell(
                    row=row_number,
                    column=status_column,
                ).value
                or ""
            ).upper()

            if current_status == "SENT":
                raise PermissionError(
                    f"Outreach {run_id} has "
                    "already been sent."
                )

            if current_status != "SENDING":
                raise PermissionError(
                    f"Outreach {run_id} must be "
                    f"SENDING before marking SENT. "
                    f"Current status: {current_status}"
                )
            if worksheet.cell(row=row_number, column=retry_count_column).value is None:
                worksheet.cell(
                    row=row_number,
                    column=retry_count_column,
                ).value = 0
            created_at = (
                __import__("datetime")
                .datetime.now()
                .astimezone()
                .isoformat()
            )

            worksheet.cell(
                row=row_number,
                column=status_column,
            ).value = "SENT"

            worksheet.cell(
                row=row_number,
                column=sent_at_column,
            ).value = created_at

            worksheet.cell(
                row=row_number,
                column=message_id_column,
            ).value = gmail_message_id

            worksheet.cell(
                row=row_number,
                column=error_column,
            ).value = None

            workbook.save(
                self.output_path
            )

            return {
                "run_id": run_id,
                "status": "SENT",
                "gmail_message_id": gmail_message_id,
                "sent_at": created_at,
            }

        raise ValueError(
            f"Outreach record not found: {run_id}"
        )

    def mark_send_failed(
        self,
        run_id: str,
        error_message: str,
    ) -> dict[str, Any]:

        if not self.output_path.exists():
            raise FileNotFoundError(
                f"Outreach tracker not found: "
                f"{self.output_path}"
            )

        workbook = load_workbook(
            self.output_path
        )

        worksheet = workbook.active

        headers = [
            cell.value
            for cell in worksheet[1]
        ]

        run_id_column = (
            headers.index("Run ID") + 1
        )

        status_column = (
            headers.index("Status") + 1
        )

        error_column = (
            headers.index("Error Message") + 1
        )
        retry_column = (
            headers.index("Retry Count") + 1
        )

        for row_number in range(
            2,
            worksheet.max_row + 1,
        ):

            current_run_id = worksheet.cell(
                row=row_number,
                column=run_id_column,
            ).value

            if current_run_id != run_id:
                continue

            current_retry_count = (
                worksheet.cell(
                    row=row_number,
                    column=retry_column,
                ).value
                or 0
            )

            worksheet.cell(
                row=row_number,
                column=status_column,
            ).value = "SEND_FAILED"

            worksheet.cell(
                row=row_number,
                column=error_column,
            ).value = error_message

            worksheet.cell(
                row=row_number,
                column=retry_column,
            ).value = (
                int(current_retry_count) + 1
            )

            workbook.save(
                self.output_path
            )

            return {
                "run_id": run_id,
                "status": "SEND_FAILED",
                "error_message": error_message,
                "retry_count": (
                    int(current_retry_count) + 1
                ),
            }

        raise ValueError(
            f"Outreach record not found: {run_id}"
        )
    def mark_retrying(
        self,
        run_id: str,
    ) -> dict[str, Any]:

        if not self.output_path.exists():
            raise FileNotFoundError(
                f"Outreach tracker not found: "
                f"{self.output_path}"
            )

        workbook = load_workbook(
            self.output_path
        )

        worksheet = workbook.active

        headers = [
            cell.value
            for cell in worksheet[1]
        ]

        run_id_column = (
            headers.index("Run ID") + 1
        )

        status_column = (
            headers.index("Status") + 1
        )

        error_column = (
            headers.index("Error Message") + 1
        )

        for row_number in range(
            2,
            worksheet.max_row + 1,
        ):

            current_run_id = worksheet.cell(
                row=row_number,
                column=run_id_column,
            ).value

            if current_run_id != run_id:
                continue

            current_status = (
                worksheet.cell(
                    row=row_number,
                    column=status_column,
                ).value
                or ""
            ).upper()

            if current_status != "SEND_FAILED":
                raise PermissionError(
                    f"Outreach {run_id} must be "
                    f"SEND_FAILED before retrying. "
                    f"Current status: {current_status}"
                )

            worksheet.cell(
                row=row_number,
                column=status_column,
            ).value = "SENDING"

            worksheet.cell(
                row=row_number,
                column=error_column,
            ).value = None

            workbook.save(
                self.output_path
            )

            return {
                "run_id": run_id,
                "previous_status": current_status,
                "status": "SENDING",
            }

        raise ValueError(
            f"Outreach record not found: {run_id}"
        )
    def record(
        self,
        message: Any,
        job: dict[str, Any],
        contact: dict[str, Any],
    ) -> dict[str, Any]:

        self.output_path.parent.mkdir(
            parents=True,
            exist_ok=True,
        )

        if self.output_path.exists():

            workbook = load_workbook(
                self.output_path
            )

            worksheet = workbook.active

            self._ensure_columns(
                worksheet
            )
        else:
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = "Outreach Tracker"

            worksheet.append(
                self.HEADERS
            )

        run_id = self._next_run_id(
            worksheet
        )

        created_at = (
            __import__("datetime")
            .datetime.now()
            .astimezone()
            .isoformat()
        )

        row = [
            run_id,
            message.job_id,
            message.company,
            message.role,
            job.get("location", ""),
            job.get("source", ""),
            message.contact_name,
            message.contact_email,
            contact.get("position", ""),
            self._get_contact_function(
                contact
            ),
            message.subject,
            message.body,
            message.cv_path,
            message.status,
            created_at,
            None,   # Sent At
            None,   # Gmail Message ID
            None,   # Error Message
            0,      # Retry Count
        ]

        worksheet.append(row)

        workbook.save(
            self.output_path
        )

        return {
            "run_id": run_id,
            "path": str(
                self.output_path
            ),
            "status": message.status,
        }
    def get_outreach(
        self,
        run_id: str,
    ) -> dict[str, Any] | None:

        if not self.output_path.exists():
            return None

        workbook = load_workbook(
            self.output_path
        )

        worksheet = workbook.active

        headers = [
            cell.value
            for cell in worksheet[1]
        ]

        for row in worksheet.iter_rows(
            min_row=2,
            values_only=True,
        ):

            record = dict(
                zip(
                    headers,
                    row,
                )
            )

            if record.get("Run ID") == run_id:
                return record

        return None

    def update_status(
        self,
        run_id: str,
        status: str,
    ) -> dict[str, Any]:

        if not self.output_path.exists():
            raise FileNotFoundError(
                f"Outreach tracker not found: "
                f"{self.output_path}"
            )

        workbook = load_workbook(
            self.output_path
        )

        worksheet = workbook.active

        headers = [
            cell.value
            for cell in worksheet[1]
        ]

        if "Status" not in headers:
            raise ValueError(
                "Outreach tracker does not contain "
                "a Status column."
            )

        status_column = (
            headers.index("Status") + 1
        )

        run_id_column = (
            headers.index("Run ID") + 1
        )

        for row_number in range(
            2,
            worksheet.max_row + 1,
        ):

            current_run_id = worksheet.cell(
                row=row_number,
                column=run_id_column,
            ).value

            if current_run_id != run_id:
                continue

            current_status = worksheet.cell(
                row=row_number,
                column=status_column,
            ).value

            if current_status == "SENT":
                raise PermissionError(
                    f"Outreach {run_id} has "
                    "already been sent."
                )

            worksheet.cell(
                row=row_number,
                column=status_column,
            ).value = status

            workbook.save(
                self.output_path
            )

            return {
                "run_id": run_id,
                "previous_status": current_status,
                "status": status,
                "path": str(
                    self.output_path
                ),
            }

        raise ValueError(
            f"Outreach record not found: {run_id}"
        )
    @staticmethod
    def _next_run_id(
        worksheet,
    ) -> str:

        existing_count = (
            worksheet.max_row - 1
        )

        return (
            f"OUT-{existing_count + 1:06d}"
        )

    @staticmethod
    def _get_contact_function(
        contact: dict[str, Any],
    ) -> str:

        position = (
            contact.get("position")
            or ""
        ).lower()

        if any(
            keyword in position
            for keyword in [
                "talent acquisition",
                "recruiter",
                "recruiting",
                "recruitment",
            ]
        ):
            return "Talent Acquisition"

        if any(
            keyword in position
            for keyword in [
                "business analysis",
                "business analyst",
            ]
        ):
            return "Business Analysis"

        if any(
            keyword in position
            for keyword in [
                "business transformation",
                "digital transformation",
                "transformation",
            ]
        ):
            return "Business Transformation"

        if any(
            keyword in position
            for keyword in [
                "product manager",
                "product director",
                "head of product",
            ]
        ):
            return "Product"

        if any(
            keyword in position
            for keyword in [
                "technology",
                "engineering",
                "architecture",
            ]
        ):
            return "Technology"

        if "digital" in position:
            return "Digital Transformation"

        return contact.get(
            "position",
            "",
        )