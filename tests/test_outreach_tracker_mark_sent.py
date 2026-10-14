from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import Workbook, load_workbook
from app.networking.outreach_tracker import OutreachTracker


class TestMarkSent(unittest.TestCase):
    def setUp(self):
        self.temp_dir = TemporaryDirectory()
        self.path = Path(self.temp_dir.name) / "tracker.xlsx"

        workbook = Workbook()
        sheet = workbook.active
        sheet.append(OutreachTracker.HEADERS)
        row = {
            "Run ID": "OUT-TEST-001",
            "Job ID": 42,
            "Status": "SENDING",
            "Retry Count": 0,
        }
        sheet.append([row.get(h) for h in OutreachTracker.HEADERS])
        workbook.save(self.path)
        workbook.close()
        self.tracker = OutreachTracker(output_path=self.path)

    def tearDown(self):
        self.temp_dir.cleanup()

    def test_mark_sent_persists_status_and_message_id(self):
        result = self.tracker.mark_sent(
            "OUT-TEST-001", "gmail-test-message-123"
        )
        self.assertEqual(result["status"], "SENT")
        self.assertEqual(
            result["gmail_message_id"], "gmail-test-message-123"
        )

        workbook = load_workbook(self.path, data_only=True)
        try:
            sheet = workbook.active
            headers = [cell.value for cell in sheet[1]]
            values = dict(zip(headers, [cell.value for cell in sheet[2]]))
        finally:
            workbook.close()

        self.assertEqual(values["Status"], "SENT")
        self.assertEqual(
            values["Gmail Message ID"], "gmail-test-message-123"
        )
        self.assertIsNotNone(values["Sent At"])
        self.assertIsNone(values["Error Message"])

    def test_mark_sent_rejects_non_sending_status(self):
        workbook = load_workbook(self.path)
        try:
            sheet = workbook.active
            headers = [cell.value for cell in sheet[1]]
            sheet.cell(
                row=2, column=headers.index("Status") + 1
            ).value = "APPROVED"
            workbook.save(self.path)
        finally:
            workbook.close()

        with self.assertRaises(PermissionError):
            self.tracker.mark_sent(
                "OUT-TEST-001", "gmail-test-message-123"
            )


if __name__ == "__main__":
    unittest.main()
